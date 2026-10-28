"""Regression tests for the 7285bcf review. Model calls are mocked."""

from __future__ import annotations

import csv
import io
import json
from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase

from erpnext_ai_assistant.api.chat import _approve_import, reject_action
from erpnext_ai_assistant.services.imports import corrected_table, preview_import, propose_import
from erpnext_ai_assistant.services.rag.service import backfill_content_norm, retrieve
from erpnext_ai_assistant.services.sheets import _export_value, _save_csv, _write_literal_row
from erpnext_ai_assistant.services.targeting import _covered_day_filters
from erpnext_ai_assistant.services.tools.agent import NOT_VERIFIED_EN, _finish_with_evidence, run_tool_loop
from erpnext_ai_assistant.services.writes import payload_hash


def _request(method="POST"):
	frappe.local.request = type("Req", (), {"method": method})()


class TestReviewFixes(FrappeTestCase):
	def setUp(self):
		super().setUp()
		frappe.db.commit = lambda *args, **kwargs: None
		_request()

	def test_leave_filter_does_not_read_protected_rows(self):
		def forbidden(*args, **kwargs):
			raise AssertionError("protected leave rows were read")

		meta = frappe._dict(
			get_field=lambda name: frappe._dict(fieldtype="Link", options="Employee", fieldname=name)
			if name == "employee"
			else None
		)
		with patch("erpnext_ai_assistant.services.targeting._span_doctypes", return_value=[{"doctype": "Leave Application", "link": "employee"}]):
			with patch("erpnext_ai_assistant.services.targeting._span_matches", return_value=True):
				with patch("erpnext_ai_assistant.services.targeting.frappe.get_meta", return_value=meta):
					with patch("erpnext_ai_assistant.services.targeting.frappe.db.exists", return_value=True):
						with patch("erpnext_ai_assistant.services.targeting.frappe.has_permission", return_value=False):
							with patch("erpnext_ai_assistant.services.targeting.frappe.get_all", side_effect=forbidden):
								filters, _note, problem = _covered_day_filters(
									"Attendance",
									{"date": "attendance_date", "group": "employee"},
									"exclude leave days",
									[["Attendance", "attendance_date", "between", ["2026-01-01", "2026-01-31"]]],
									"en",
								)
		self.assertEqual(filters, [])
		self.assertIn("cannot be verified", problem)
		self.assertIn("Leave Application", problem)

	def test_leave_filter_respects_hidden_dates(self):
		class Meta:
			def get_field(self, name):
				kind = "Link" if name == "employee" else "Date"
				return frappe._dict(fieldtype=kind, fieldname=name, options="Employee" if name == "employee" else None)

			def get_permitted_fieldnames(self, parenttype=None):
				return ["employee", "status"]

		with patch("erpnext_ai_assistant.services.targeting._span_doctypes", return_value=[{"doctype": "Leave Application", "link": "employee"}]):
			with patch("erpnext_ai_assistant.services.targeting._span_matches", return_value=True):
				with patch("erpnext_ai_assistant.services.targeting.frappe.get_meta", return_value=Meta()):
					with patch("erpnext_ai_assistant.services.targeting.frappe.db.exists", return_value=True):
						with patch("erpnext_ai_assistant.services.targeting.frappe.has_permission", return_value=True):
							with patch("erpnext_ai_assistant.services.targeting.frappe.get_all", side_effect=AssertionError("field bypass")):
								_filters, _note, problem = _covered_day_filters(
									"Attendance",
									{"date": "attendance_date", "group": "employee"},
									"do not count leave days",
									[["Attendance", "attendance_date", "between", ["2026-01-01", "2026-01-31"]]],
									"en",
								)
		self.assertIn("from_date", problem)
		self.assertIn("permission level", problem)

	def test_reject_is_whitelisted_and_helper_is_not(self):
		self.assertIn(reject_action, frappe.whitelisted)
		self.assertNotIn(_approve_import, frappe.whitelisted)
		from frappe.handler import execute_cmd

		frappe.set_user("Guest")
		frappe.local.form_dict = frappe._dict(cmd="erpnext_ai_assistant.api.chat._approve_import")
		with self.assertRaises(frappe.PermissionError):
			execute_cmd("erpnext_ai_assistant.api.chat._approve_import")

	def test_http_reject_owner_and_repeat(self):
		from frappe.handler import execute_cmd

		frappe.set_user("Administrator")
		doc = frappe.get_doc(
			{
				"doctype": "ERP AI Pending Action",
				"title": "reject-" + frappe.generate_hash(length=8),
				"user": "Administrator",
				"action_type": "update",
				"target_doctype": "ToDo",
				"payload_json": "{}",
				"payload_hash": payload_hash({}),
				"summary": "reject",
				"status": "pending",
			}
		).insert()
		frappe.local.form_dict = frappe._dict(
			cmd="erpnext_ai_assistant.api.chat.reject_action",
			pending_action=doc.name,
			reason="no",
		)
		result = execute_cmd("erpnext_ai_assistant.api.chat.reject_action")
		self.assertEqual(result["status"], "rejected")
		with self.assertRaises(frappe.ValidationError):
			execute_cmd("erpnext_ai_assistant.api.chat.reject_action")

	def test_live_answer_accepted_after_count(self):
		calls = {"n": 0}

		def chat_message(messages, tools=None, **kwargs):
			calls["n"] += 1
			if calls["n"] == 1:
				return {
					"content": "",
					"tool_calls": [
						{"id": "c1", "function": {"name": "count_records", "arguments": json.dumps({"doctype": "ToDo"})}}
					],
				}
			return {"content": "There are 3 tasks.", "tool_calls": []}

		def dispatch(name, args, chat_id, lang, user_text=None):
			if name == "count_records":
				return {"doctype": "ToDo", "total": 3}
			return {"error": "failed"}

		with patch("erpnext_ai_assistant.services.tools.agent.llm.chat_message", side_effect=chat_message):
			with patch("erpnext_ai_assistant.services.tools.agent.llm.begin_turn_deadline", lambda: None):
				with patch("erpnext_ai_assistant.services.tools.agent.dispatch", side_effect=dispatch):
					with patch("erpnext_ai_assistant.services.tools.agent.active_tool_specs", return_value=[{"type": "function"}]):
						out = run_tool_loop("sys", "how many todos", [], "", None, None, "en", lambda *a, **k: None)
		self.assertEqual(out["answer"], "There are 3 tasks.")

	def test_live_claim_without_evidence_is_refused(self):
		def chat_message(messages, tools=None, **kwargs):
			return {"content": "There are 3 tasks.", "tool_calls": []}

		with patch("erpnext_ai_assistant.services.tools.agent.llm.chat_message", side_effect=chat_message):
			with patch("erpnext_ai_assistant.services.tools.agent.llm.begin_turn_deadline", lambda: None):
				with patch("erpnext_ai_assistant.services.tools.agent.active_tool_specs", return_value=[]):
					out = run_tool_loop("sys", "how many todos", [], "", None, None, "en", lambda *a, **k: None)
		self.assertEqual(out["answer"], NOT_VERIFIED_EN)

	def test_failed_or_unrelated_tool_is_not_evidence(self):
		refused = _finish_with_evidence(
			"There are 3 tasks.",
			"how many todos",
			"en",
			None,
			[{"kind": "had_data", "action": "search_knowledge"}],
		)
		self.assertEqual(refused["answer"], NOT_VERIFIED_EN)
		failed = _finish_with_evidence(
			"There are 3 tasks.",
			"how many todos",
			"en",
			None,
			[{"kind": "tool_error", "action": "count_records"}],
		)
		self.assertEqual(failed["answer"], NOT_VERIFIED_EN)

	def test_round_limit_still_requires_evidence(self):
		calls = {"n": 0}

		def chat_message(messages, tools=None, **kwargs):
			calls["n"] += 1
			if tools:
				return {
					"content": "",
					"tool_calls": [
						{
							"id": f"c{calls['n']}",
							"function": {"name": "describe_doctype", "arguments": json.dumps({"doctype": "ToDo"})},
						}
					],
				}
			return {"content": "There are 9 tasks.", "tool_calls": []}

		with patch("erpnext_ai_assistant.services.tools.agent.llm.chat_message", side_effect=chat_message):
			with patch("erpnext_ai_assistant.services.tools.agent.llm.begin_turn_deadline", lambda: None):
				with patch(
					"erpnext_ai_assistant.services.tools.agent.dispatch",
					return_value={"doctype": "ToDo", "fields": []},
				):
					with patch("erpnext_ai_assistant.services.tools.agent.active_tool_specs", return_value=[{"type": "function"}]):
						out = run_tool_loop("sys", "how many todos", [], "", None, None, "en", lambda *a, **k: None)
		self.assertEqual(out["answer"], NOT_VERIFIED_EN)

	def test_update_import_keeps_id_and_previews_update(self):
		text, rows = corrected_table(
			["ID", "Description"],
			[["EXISTING-TODO", "Updated task"]],
			"ToDo",
		)
		self.assertIn("ID", text.splitlines()[0])
		self.assertIn("EXISTING-TODO", rows[0])
		seen = {}

		class FakeBook:
			def __init__(self, doctype, file_url, import_type=None):
				seen["import_type"] = import_type

			def get_data_for_import_preview(self):
				return {"warnings": [], "columns": [], "total_number_of_rows": 1}

		with patch("frappe.core.doctype.data_import.importer.ImportFile", FakeBook):
			preview_import("ToDo", "/private/files/example.csv", "Update Existing Records")
		self.assertEqual(seen["import_type"], "Update Existing Records")

	def test_update_without_id_is_rejected(self):
		class FileDoc:
			file_name = "rows.csv"
			file_url = "/private/files/rows.csv"

			def get_content(self):
				return b"Description\nHello\n"

		with patch("erpnext_ai_assistant.services.imports.owned_import_file", return_value=FileDoc()):
			with patch("erpnext_ai_assistant.services.imports.import_block_reason", return_value=None):
				result = propose_import("", "/private/files/rows.csv", "ToDo", "Update Existing Records", "en")
		self.assertIsNone(result["pending"])
		self.assertIn("ID", result["message"])

	def test_import_startup_failure_can_retry_same_import(self):
		frappe.set_user("Administrator")
		payload = {
			"reference_doctype": "ToDo",
			"import_type": "Insert New Records",
			"file_url": "/private/files/todo.csv",
			"row_count": 1,
		}
		doc = frappe.get_doc(
			{
				"doctype": "ERP AI Pending Action",
				"title": "import-" + frappe.generate_hash(length=8),
				"user": "Administrator",
				"action_type": "import",
				"target_doctype": "ToDo",
				"payload_json": json.dumps(payload),
				"payload_hash": payload_hash(payload),
				"summary": "import",
				"status": "pending",
			}
		).insert()
		class DummyImport:
			name = "DI-REVIEW-1"
			status = "Pending"
			template_warnings = ""
			starts = 0

			def start_import(self):
				self.starts += 1
				if self.starts == 1:
					raise RuntimeError("queue down")
				return True

		dummy = DummyImport()
		real_get_doc = frappe.get_doc

		def get_doc(*args, **kwargs):
			if args and args[0] == "Data Import":
				return dummy
			return real_get_doc(*args, **kwargs)

		with patch("erpnext_ai_assistant.services.imports.execute_import", return_value={"data_import": dummy.name, "doctype": "ToDo"}):
			with patch("erpnext_ai_assistant.api.chat.frappe.get_doc", side_effect=get_doc):
				with self.assertRaises(frappe.ValidationError):
					_approve_import(doc, "I approve this import")
		doc.reload()
		self.assertEqual(doc.status, "approved")
		self.assertEqual(json.loads(doc.result_json)["data_import"], dummy.name)
		doc.reload()
		with patch("erpnext_ai_assistant.services.imports.execute_import", side_effect=AssertionError("second import")):
			with patch("erpnext_ai_assistant.api.chat.frappe.get_doc", side_effect=get_doc):
				with patch(
					"erpnext_ai_assistant.services.imports.import_status_message",
					return_value={"message": "ok", "files": [], "status": "Pending"},
				):
					result = _approve_import(doc, "I approve this import")
		self.assertEqual(result["status"], "executed")
		self.assertEqual(dummy.starts, 2)
		self.assertEqual(frappe.db.get_value("ERP AI Pending Action", doc.name, "status"), "executed")

	def test_retrieve_filters_normalized_text(self):
		seen = {}

		def get_all(doctype, *args, **kwargs):
			if kwargs.get("or_filters"):
				seen["or_filters"] = kwargs.get("or_filters")
			if doctype == "ERP AI Knowledge Chunk":
				return [
					frappe._dict(
						name="C1",
						document="D1",
						content="مهمة جديدة",
						chunk_index=0,
					)
				]
			if doctype == "ERP AI Knowledge Document":
				return [frappe._dict(name="D1", source_roles="")]
			return []

		with patch("erpnext_ai_assistant.services.rag.service.frappe.db.has_column", return_value=True):
			with patch("erpnext_ai_assistant.services.rag.service.frappe.get_all", side_effect=get_all):
				with patch("erpnext_ai_assistant.services.rag.service.frappe.db.get_value", return_value={"title": "Note", "source_key": "policy:1"}):
					hits = retrieve("مهمة", limit=4)
		self.assertEqual(seen["or_filters"][0][0], "content_norm")
		self.assertIn("مهمه", seen["or_filters"][0][2])
		self.assertEqual(hits[0]["document"], "D1")

	def test_arabic_retrieval_uses_normalized_column(self):
		if not frappe.db.has_column("ERP AI Knowledge Chunk", "content_norm"):
			self.skipTest("content_norm is added by migrate")
		doc = frappe.get_doc(
			{
				"doctype": "ERP AI Knowledge Document",
				"title": "Arabic task " + frappe.generate_hash(length=6),
				"content": "مهمة جديدة",
				"source_type": "policy",
				"source_roles": "",
			}
		).insert(ignore_permissions=True)
		chunk = frappe.get_doc(
			{
				"doctype": "ERP AI Knowledge Chunk",
				"title": "Arabic chunk " + frappe.generate_hash(length=6),
				"document": doc.name,
				"chunk_index": 0,
				"content": "مهمة جديدة",
			}
		).insert(ignore_permissions=True)
		backfill_content_norm()
		chunk.reload()
		self.assertIn("مهمه", chunk.content_norm)
		hits = retrieve("مهمة", limit=4)
		self.assertTrue(any(hit.get("document") == doc.name for hit in hits))

	def test_export_text_is_not_a_formula(self):
		import openpyxl

		book = openpyxl.Workbook()
		sheet = book.active
		_write_literal_row(
			sheet,
			1,
			[_export_value("=1+1"), _export_value("+1"), _export_value(3), _export_value("مهمة"), _export_value(" =1+1")],
		)
		# XLSX: a string cell is never evaluated, so the text stays exactly as stored.
		self.assertEqual(sheet.cell(1, 1).data_type, "s")
		self.assertEqual(sheet.cell(1, 1).value, "=1+1")
		self.assertEqual(sheet.cell(1, 3).value, 3)
		self.assertEqual(sheet.cell(1, 4).value, "مهمة")
		saved = {}
		with patch("erpnext_ai_assistant.services.sheets._store", side_effect=lambda name, content, label: saved.update(content=content)):
			_save_csv("rows.csv", ["A", "B", "C", "D"], [[_export_value("=1+1"), _export_value("@SUM(A1)"), _export_value(-5), _export_value("-5")]], "en")
		row = list(csv.reader(io.StringIO(saved["content"].decode("utf-8-sig"))))[1]
		self.assertEqual(row[0], "'=1+1")
		self.assertEqual(row[1], "'@SUM(A1)")
		self.assertEqual(row[2], "-5")
		self.assertEqual(row[3], "'-5")
